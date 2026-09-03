#!/usr/bin/env python3
"""Generate feature list docx for both Track A and Track B."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
import csv
from collections import OrderedDict

PROJ = Path(__file__).resolve().parent.parent
OUT = PROJ / "TrackA_Features.docx"
doc = Document()

NAVY = RGBColor.from_string("1F4E79")

h = doc.add_heading("特征完整清单", level=1)
for r in h.runs:
    r.font.color.rgb = NAVY

# Data ranges
doc.add_heading("数据覆盖范围", level=2)
doc.add_paragraph("Track A（E0-E3 四级联赛）：2019-20 ~ 2024-25，共 6 赛季 11,952 场", style="List Bullet")
doc.add_paragraph("Track B（英超 + FootyStats）：2019-20 ~ 2024-25，共 6 赛季 2,280 场", style="List Bullet")
doc.add_paragraph("附注：所有赔率列（odds_home, odds_draw 等 41 列）均不参与训练，仅作为后验对比基准。", style="List Bullet")

# Read features
with open(PROJ / "trackA_features.csv", 'r', encoding='utf-8') as f:
    reader = csv.DictReader(f)
    rows = list(reader)

# Group by track then category
tracks = OrderedDict()
for r in rows:
    track = r['track']
    cat = r['category']
    tracks.setdefault(track, OrderedDict())
    tracks[track].setdefault(cat, [])
    tracks[track][cat].append(r['feature_name'])

for track, categories in tracks.items():
    if track == 'A':
        doc.add_heading(f"Track A：62 维（基础特征）", level=2)
    else:
        doc.add_heading(f"Track B：+9 维（仅当有 FootyStats 数据时）", level=2)

    total = 0
    for cat, feats in categories.items():
        doc.add_heading(f"{cat}（{len(feats)} 个）", level=3)
        for f in feats:
            doc.add_paragraph(f, style="List Bullet")
        total += len(feats)

    if track == 'A':
        p = doc.add_paragraph()
        r = p.add_run(f"Track A 合计：{total} 维")
        r.bold = True
    else:
        p = doc.add_paragraph()
        r = p.add_run(f"Track B 附加合计：{total} 维（总 62+9=71 维）")
        r.bold = True

p = doc.add_paragraph()
r = p.add_run(f"\n总特征集：62 维（Track A）+ 9 维（Track B FootyStats）= 71 维")
r.bold = True
r.font.size = Pt(12)

doc.save(OUT)
print(f"Saved: {OUT}")
